"""Build Chapter IV DOCX and insert it into a copy of the group template."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs/06_report/BAB_IV_UAS_KSI_KELOMPOK_2.md"
SOURCE = Path(r"C:\Users\daffa\Documents\Laporan Proyek Akhir KSI_Kelompok 2.docx")
OUT = MD.parent


def style_doc(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        if name == "Normal":
            s.font.size = Pt(12)
            s.paragraph_format.line_spacing = 1.5
            s.paragraph_format.space_after = Pt(6)
        elif name == "Heading 1": s.font.size = Pt(14)
        elif name == "Heading 2": s.font.size = Pt(13)
        elif name == "Heading 3": s.font.size = Pt(12)


def rich(p, line):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", line)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        else:
            p.add_run(part)


def add_table(doc, lines):
    rows = [re.split(r"\s*\|\s*", s.strip().strip("|")) for s in lines]
    rows = [rows[0], *rows[2:]]  # separator line
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.text = value.replace("**", "").replace("`", "")
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5 if len(rows[0]) > 5 else 9)
                    if i == 0: run.bold = True
    doc.add_paragraph()


def add_markdown(doc):
    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                buf.append(lines[i]); i += 1
            add_table(doc, buf); continue
        if line.startswith("!["):
            m = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if m:
                p = MD.parent / m.group(1)
                if not p.exists(): raise FileNotFoundError(p)
                para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.add_run().add_picture(str(p), width=Cm(15.5))
            i += 1; continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("*Gambar "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.strip("*")); r.italic = True; r.font.size = Pt(9)
        else:
            p = doc.add_paragraph()
            rich(p, line)
        i += 1


def main():
    standalone = Document()
    style_doc(standalone)
    add_markdown(standalone)
    chapter = OUT / "BAB_IV_UAS_KSI_KELOMPOK_2_DOKUMENTASI_RUN.docx"
    standalone.save(chapter)

    doc = Document(SOURCE)
    style_doc(doc)
    body = doc.element.body
    children = list(body)
    def text(el):
        return "".join(node.text or "" for node in el.iter() if node.tag.endswith("}t"))
    begin = next(i for i, el in enumerate(children) if text(el).startswith("BAB IV — HASIL IMPLEMENTASI DAN PENGUJIAN"))
    end = next(i for i, el in enumerate(children) if i > begin and text(el).startswith("BAB V — RANCANGAN PENGELOLAAN"))
    for el in children[begin:end]: body.remove(el)
    insert_at = begin
    old_end = len(body) - 1  # section properties remain last
    add_markdown(doc)
    new = list(body)[old_end:-1]
    for index, el in enumerate(new):
        body.remove(el)
        body.insert(insert_at + index, el)
    integrated = OUT / "Laporan_Proyek_Akhir_KSI_Kelompok_2_Bab_IV_DOKUMENTASI_RUN.docx"
    doc.save(integrated)
    print(chapter)
    print(integrated)


if __name__ == "__main__":
    main()
